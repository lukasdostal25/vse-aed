"""Export whole Excel workbooks (values, formulas, fills, colour scales) to encrypted JSON for the in-app preview.

  ECON_PW=... python3 xlbook.py APP OUT_REL XLSX [MAX_ROWS]
  OUT_REL e.g. s01/wb_dta1_indpro_solved  ->  APP/data/decks/s01/wb_dta1_indpro_solved.json
"""
import base64, datetime as dt, json, os, re, sys
import openpyxl
from openpyxl.utils import range_boundaries
sys.path.insert(0, os.path.dirname(__file__))
from deck import key_from, enc

def num(v, nf):
    if isinstance(v, bool):
        return 'PRAVDA' if v else 'NEPRAVDA'
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime('%d.%m.%Y')
    if isinstance(v, (int, float)):
        if '%' in (nf or ''):
            d = len(nf.split('.')[1].rstrip('%')) if '.' in nf else 0
            return ('%.*f %%' % (d, v * 100)).replace('.', ',')
        if isinstance(v, int) or float(v).is_integer():
            return str(int(v))
        m = re.search(r'0\.(0+)', nf or '')
        s = ('%.*f' % (len(m.group(1)), v)) if m else ('%.6g' % v if abs(v) < 1e6 else '%.0f' % v)
        return s.replace('.', ',')
    return str(v)

def cz_formula(f):
    # Czech Excel: ';' separates arguments, ',' is the decimal mark
    out, q = [], False
    for ch in f:
        if ch == '"':
            q = not q
        if not q and ch == ',':
            ch = ';'
        out.append(ch)
    s = ''.join(out)
    s = re.sub(r'(?<=\d)\.(?=\d)', ',', s)
    for en, cz in CZ_FUN.items():
        s = re.sub(r'\b%s\(' % en, cz + '(', s)
    return s

# function names as they appear in the Czech Excel
CZ_FUN = {'AVERAGE': 'PRŮMĚR', 'SQRT': 'ODMOCNINA', 'SUM': 'SUMA', 'IF': 'KDYŽ', 'COUNT': 'POČET', 'ROUND': 'ZAOKROUHLIT',
          'AVERAGEIF': 'PRŮMĚRKDYŽ', 'COUNTIF': 'COUNTIF', 'ABS': 'ABS', 'POWER': 'POWER'}

def hexcol(c):
    try:
        r = c.rgb
        if isinstance(r, str) and len(r) == 8 and r != '00000000':
            return '#' + r[2:]
    except Exception:
        pass
    return None

def mix(c1, c2, t):
    a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]; b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
    return '#%02x%02x%02x' % tuple(round(x + (y - x) * t) for x, y in zip(a, b))

def export(app, out_rel, path, max_rows=400, max_cols=40):
    wf = openpyxl.load_workbook(path)                 # formulas
    wv = openpyxl.load_workbook(path, data_only=True) # cached values
    sheets = []
    for ws in wf:
        vs = wv[ws.title]
        nr, nc = min(ws.max_row, max_rows), min(ws.max_column, max_cols)
        bg = {}
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                if rule.type != 'colorScale' or not rule.colorScale:
                    continue
                cols = [hexcol(c) or '#ffffff' for c in rule.colorScale.color]
                for rng in str(cf.sqref).split():
                    c0, r0, c1, r1 = range_boundaries(rng)
                    cells = [(r, c) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)]
                    vals = [vs.cell(r, c).value for r, c in cells]
                    nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
                    if not nums:
                        continue
                    lo, hi = min(nums), max(nums)
                    for (r, c), v in zip(cells, vals):
                        if not isinstance(v, (int, float)) or isinstance(v, bool):
                            continue
                        t = 0 if hi == lo else (v - lo) / (hi - lo)
                        bg[(r, c)] = mix(cols[0], cols[-1], t) if len(cols) == 2 else (mix(cols[0], cols[1], t * 2) if t <= .5 else mix(cols[1], cols[2], t * 2 - 1))
        rows = []
        for r in range(1, nr + 1):
            row = []
            for c in range(1, nc + 1):
                cf_, cv = ws.cell(r, c), vs.cell(r, c)
                f = cf_.value if isinstance(cf_.value, str) and cf_.value.startswith('=') else None
                v = cv.value
                if v is None and f is None:
                    row.append(None); continue
                o = {'v': '' if v is None else num(v, cv.number_format)}
                if isinstance(v, str):
                    o['t'] = 1
                if f:
                    o['f'] = cz_formula(f)
                col = bg.get((r, c))
                if not col and cv.fill and cv.fill.fill_type == 'solid':
                    col = hexcol(cv.fill.fgColor)
                if col and col.lower() not in ('#ffffff',):
                    o['bg'] = col
                row.append(o if len(o) > 1 or 't' in o else o['v'])
            while row and row[-1] is None:
                row.pop()
            rows.append(row)
        sh = {'name': ws.title, 'rows': rows}
        notes = []
        if ws.max_row > max_rows:
            notes.append('Zobrazeno prvních %d z %d řádků.' % (max_rows, ws.max_row))
        if ws.max_column > max_cols:
            notes.append('Zobrazeno prvních %d z %d sloupců.' % (max_cols, ws.max_column))
        if getattr(ws, '_charts', None):
            notes.append('List obsahuje %d graf(y) – viz výsledné grafy u cvičení.' % len(ws._charts))
        if notes:
            sh['note'] = ' '.join(notes)
        sheets.append(sh)
    data = {'name': os.path.basename(path), 'sheets': sheets}
    key = key_from(app)
    iv, ct = enc(key, json.dumps(data, ensure_ascii=False, separators=(',', ':')).encode())
    out = os.path.join(app, 'data', 'decks', out_rel + '.json')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    json.dump({'iv': base64.b64encode(iv).decode(), 'ct': base64.b64encode(ct).decode()}, open(out, 'w'))
    print(out_rel, [(s['name'], len(s['rows'])) for s in sheets], os.path.getsize(out) // 1024, 'kB')

if __name__ == '__main__':
    export(sys.argv[1], sys.argv[2], sys.argv[3], int(sys.argv[4]) if len(sys.argv) > 4 else 400)
